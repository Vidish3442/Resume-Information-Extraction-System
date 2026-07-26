import pytest
from pathlib import Path
from docx import Document


@pytest.fixture
def skills_file(tmp_path):
    """Write a temporary skills.txt with a known set of entries and return its path."""
    content = "# Programming languages\npython\njavascript\nsql\nfastapi\nreact\ndocker\n# Data\nmachine learning\n"
    p = tmp_path / "skills.txt"
    p.write_text(content, encoding="utf-8")
    return str(p)


@pytest.fixture
def sample_docx(tmp_path):
    """Create a minimal DOCX with resume content and return its path."""
    doc = Document()
    doc.add_paragraph("John Smith")
    doc.add_paragraph("john.smith@email.com")
    doc.add_paragraph("+91 9876543210")
    doc.add_paragraph("linkedin.com/in/johnsmith")
    doc.add_paragraph("github.com/johnsmith")
    doc.add_paragraph("Skills")
    doc.add_paragraph("Python SQL FastAPI")
    doc.add_paragraph("Education")
    doc.add_paragraph("B.Tech Computer Science ABC University 2026")
    doc.add_paragraph("Experience")
    doc.add_paragraph("Software Engineer Intern at XYZ Company Jan 2025 - Present")
    path = tmp_path / "sample_resume.docx"
    doc.save(str(path))
    return path


@pytest.fixture
def sample_pdf(tmp_path):
    """
    Return path to a sample PDF. 
    If sample_resumes/ directory has a PDF, copy it. 
    Otherwise create a minimal one using reportlab if available, 
    or skip by returning None.
    """
    import shutil
    # Try to copy from sample_resumes/
    sample_dir = Path(__file__).parent.parent / "sample_resumes"
    pdfs = list(sample_dir.glob("*.pdf")) if sample_dir.exists() else []
    if pdfs:
        dest = tmp_path / pdfs[0].name
        shutil.copy(pdfs[0], dest)
        return dest
    # Try to create minimal PDF with reportlab
    try:
        from reportlab.pdfgen import canvas as rl_canvas
        path = tmp_path / "sample_resume.pdf"
        c = rl_canvas.Canvas(str(path))
        c.drawString(100, 750, "John Smith")
        c.drawString(100, 730, "john.smith@email.com")
        c.drawString(100, 710, "+91 9876543210")
        c.save()
        return path
    except ImportError:
        return None
